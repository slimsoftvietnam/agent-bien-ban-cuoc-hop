"""Xuất bảng Action Items ra Excel (openpyxl, miễn phí)."""
from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

from mmlib import UNKNOWN, AgentError

STATUSES = ["Chưa bắt đầu", "Đang làm", "Hoàn thành", "Tạm dừng", "Bị chặn"]
TEAL, ORANGE, PAPER = "1F5C6B", "D9693A", "FAF6EE"


def _ev_text(it: dict[str, Any]) -> tuple[str, str]:
    evs = it.get("evidence") or []
    quote = " | ".join(f'"{e["quote"]}" — {e["speaker"]}' for e in evs)
    when = ", ".join(e["time"] or f'dòng {e.get("line")}' for e in evs)
    return quote, when


def build_xlsx(m: dict[str, Any], path: Path, only_owner: str | None = None) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        raise AgentError("Chưa cài openpyxl để xuất Excel.", "Chạy setup.ps1 hoặc: python -m pip install openpyxl")

    meeting = m.get("meeting", {})
    actions = [a for a in m.get("action_items", []) if not only_owner or a["owner"] == only_owner]
    wb = Workbook()
    ws = wb.active
    ws.title = "Action Items"
    thin = Side(style="thin", color="E4DDD0")
    border = Border(bottom=thin, top=thin, left=thin, right=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    ws["A1"] = ("Phiếu giao việc — " + only_owner) if only_owner else "Bảng theo dõi công việc sau cuộc họp"
    ws["A1"].font = Font(size=16, bold=True, color=TEAL)
    ws["A2"] = f'{meeting.get("title", "")} · Ngày họp: {meeting.get("date") or "Chưa nêu"}'
    ws["A2"].font = Font(size=11, italic=True, color="5A636B")
    headers = ["STT", "Mã", "Công việc", "Người phụ trách", "Email", "Hỗ trợ", "Thời hạn", "Thời hạn (nguyên văn)",
               "Ưu tiên", "Trạng thái", "Kết quả cần bàn giao", "Dẫn chứng", "Mốc thời gian", "Ghi chú"]
    widths = [6, 7, 46, 22, 28, 18, 13, 22, 11, 15, 30, 60, 14, 28]
    hr = 4
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=hr, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[hr].height = 30
    for i, a in enumerate(actions, 1):
        quote, when = _ev_text(a)
        dl = None
        if a.get("deadline"):
            try:
                dl = date.fromisoformat(str(a["deadline"]))
            except ValueError:
                dl = None
        row = [i, a["id"], a.get("task", ""), a["owner"], a.get("owner_email") or ("" if a["owner"] == UNKNOWN else "Cần bổ sung"),
               ", ".join(a.get("supporters") or []), dl or "Chưa nêu", a.get("deadline_text") or "",
               a.get("priority", "Trung bình"), a.get("status", "Chưa bắt đầu"), a.get("deliverable", ""), quote, when,
               a.get("notes", "")]
        r = hr + i
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = wrap
            cell.border = border
            if i % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=PAPER)
        if dl:
            ws.cell(row=r, column=7).number_format = "DD/MM/YYYY"
        if a["owner"] == UNKNOWN:
            ws.cell(row=r, column=4).font = Font(bold=True, color="B23A2B")
        if ws.cell(row=r, column=5).value == "Cần bổ sung":
            ws.cell(row=r, column=5).font = Font(italic=True, color="B23A2B")
    last = hr + max(1, len(actions))
    ws.freeze_panes = ws.cell(row=hr + 1, column=4)
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(headers))}{last}"
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"Cao,Trung bình,Thấp"', allow_blank=True)
    ws.add_data_validation(dv)
    ws.add_data_validation(dv2)
    dv.add(f"J{hr + 1}:J{last + 50}")
    dv2.add(f"I{hr + 1}:I{last + 50}")
    # quá hạn → đỏ, hoàn thành → xanh
    ws.conditional_formatting.add(f"A{hr + 1}:N{last + 50}", FormulaRule(
        formula=[f'AND(ISNUMBER($G{hr + 1}),$G{hr + 1}<TODAY(),$J{hr + 1}<>"Hoàn thành")'],
        fill=PatternFill("solid", fgColor="FBE7E4")))
    ws.conditional_formatting.add(f"A{hr + 1}:N{last + 50}", FormulaRule(
        formula=[f'$J{hr + 1}="Hoàn thành"'], fill=PatternFill("solid", fgColor="E3F0E6")))
    ws.sheet_view.zoomScale = 100

    if not only_owner:
        # tổng hợp theo người — dùng công thức để tự cập nhật khi đổi trạng thái
        s = wb.create_sheet("Theo người phụ trách")
        s["A1"] = "Tổng hợp theo người phụ trách (tự cập nhật theo sheet Action Items)"
        s["A1"].font = Font(size=14, bold=True, color=TEAL)
        hdr = ["Người phụ trách", "Tổng việc", "Hoàn thành", "Đang làm", "Chưa bắt đầu", "Quá hạn"]
        for c, h in enumerate(hdr, 1):
            cell = s.cell(row=3, column=c, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=ORANGE)
            s.column_dimensions[get_column_letter(c)].width = 24 if c == 1 else 14
        owners = sorted({a["owner"] for a in actions})
        rng = f"'Action Items'!$D${hr + 1}:$D${last + 50}"
        st = f"'Action Items'!$J${hr + 1}:$J${last + 50}"
        dlr = f"'Action Items'!$G${hr + 1}:$G${last + 50}"
        for i, o in enumerate(owners, 4):
            s.cell(row=i, column=1, value=o)
            s.cell(row=i, column=2, value=f"=COUNTIF({rng},A{i})")
            s.cell(row=i, column=3, value=f'=COUNTIFS({rng},A{i},{st},"Hoàn thành")')
            s.cell(row=i, column=4, value=f'=COUNTIFS({rng},A{i},{st},"Đang làm")')
            s.cell(row=i, column=5, value=f'=COUNTIFS({rng},A{i},{st},"Chưa bắt đầu")')
            s.cell(row=i, column=6, value=f'=COUNTIFS({rng},A{i},{dlr},"<"&TODAY(),{st},"<>Hoàn thành")')

        d = wb.create_sheet("Quyết định")
        d["A1"] = "Các quyết định đã chốt"
        d["A1"].font = Font(size=14, bold=True, color=TEAL)
        for c, (h, w) in enumerate(zip(["Mã", "Quyết định", "Lý do / bối cảnh", "Dẫn chứng", "Mốc"], [7, 50, 36, 60, 14]), 1):
            cell = d.cell(row=3, column=c, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="3C7A4F")
            d.column_dimensions[get_column_letter(c)].width = w
        for i, it in enumerate(m.get("decisions", []), 4):
            q, w = _ev_text(it)
            for c, v in enumerate([it["id"], it.get("text", ""), it.get("rationale", ""), q, w], 1):
                d.cell(row=i, column=c, value=v).alignment = wrap

    g = wb.create_sheet("Hướng dẫn")
    for i, line in enumerate([
        "Cách dùng bảng này",
        "• Cột 'Trạng thái' và 'Ưu tiên' có danh sách chọn sẵn — bấm vào ô để chọn.",
        "• Dòng tô đỏ = đã quá hạn mà chưa hoàn thành; tô xanh = đã hoàn thành.",
        "• 'Chưa nêu' = cuộc họp không nói thời hạn; 'Chưa xác định' = chưa rõ người phụ trách → cần bổ sung.",
        "• Cột 'Dẫn chứng' chép nguyên văn câu nói trong cuộc họp để đối chiếu.",
        "• Có thể tải file lên Google Drive và mở bằng Google Sheets — công thức vẫn hoạt động."], 1):
        g.cell(row=i, column=1, value=line).font = Font(bold=(i == 1), size=13 if i == 1 else 11)
    g.column_dimensions["A"].width = 110
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
