"""Tự kiểm thử end-to-end với dữ liệu mẫu + các ca lỗi, rồi xuất báo cáo PEAS (bao-cao-tu-kiem-thu.md)."""
from __future__ import annotations

import copy
import json
import shutil
import time
from collections import Counter
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from mmlib import SKILL_DIR, UNKNOWN, AgentError, log, norm_text, read_json, write_json

RESULTS: list[dict[str, Any]] = []


def check(group: str, name: str, ok: bool, detail: str = "") -> bool:
    RESULTS.append({"group": group, "name": name, "ok": bool(ok), "detail": detail})
    log(f"   {'✅' if ok else '❌'} [{group}] {name}{' — ' + detail if detail else ''}")
    return ok


def expect_error(group: str, name: str, fn, must_contain: str = "") -> None:
    try:
        fn()
        check(group, name, False, "không báo lỗi")
    except AgentError as e:
        check(group, name, must_contain.lower() in str(e).lower(), str(e)[:110])
    except SystemExit as e:
        check(group, name, True, f"dừng với mã {e.code}")


# ---------------------------------------------------------------- tái neo dẫn chứng cho transcript phiên âm
def best_substring(quote: str, segs: list[dict[str, Any]]) -> tuple[str, str, float]:
    qn = norm_text(quote).split()
    L = len(qn)
    best = ("", "", 0.0)
    for i, s in enumerate(segs):
        words = s["text"].split()
        if len(set(norm_text(s["text"]).split()) & set(qn)) < 0.3 * L:
            continue
        for size in range(max(1, L - 3), L + 4):
            for j in range(0, max(1, len(words) - size + 1)):
                cand = " ".join(words[j:j + size])
                r = SequenceMatcher(None, " ".join(qn), norm_text(cand)).ratio()
                if r > best[2]:
                    best = (cand.strip(" ,.;:?!"), s["id"], r)
    return best


def speaker_map_from_truth(out: Path, gt: list[dict[str, Any]]) -> tuple[dict[str, Any], float]:
    segs = read_json(out / "work" / "transcript.json")["segments"]
    votes: dict[str, Counter] = {}
    agree = total = 0.0
    for s in segs:
        mid = (s["start"] + s["end"]) / 2
        g = next((g for g in gt if g["start"] - 0.3 <= mid <= g["end"] + 0.3), None)
        if g:
            votes.setdefault(s["speaker"], Counter())[g["speaker"]] += s["end"] - s["start"]
    smap = {lab: {"name": c.most_common(1)[0][0], "basis": "Tự kiểm thử: đối chiếu đáp án chuẩn của dữ liệu mẫu",
                  "confidence": "cao"} for lab, c in votes.items()}
    for s in segs:
        mid = (s["start"] + s["end"]) / 2
        g = next((g for g in gt if g["start"] - 0.3 <= mid <= g["end"] + 0.3), None)
        if g:
            total += s["end"] - s["start"]
            agree += (s["end"] - s["start"]) if smap.get(s["speaker"], {}).get("name") == g["speaker"] else 0
    return smap, (agree / total if total else 0.0)


def reanchor(minutes: dict[str, Any], out: Path) -> tuple[dict[str, Any], float]:
    segs = read_json(out / "work" / "transcript.json")["segments"]
    m = copy.deepcopy(minutes)
    ratios = []
    for sec in ("summary", "timeline", "decisions", "action_items", "open_issues", "key_moments"):
        for it in m.get(sec, []):
            for ev in it.get("evidence", []):
                q, sid, r = best_substring(ev["quote"], segs)
                ratios.append(r)
                ev["quote"], ev["seg"] = q, sid
            if sec == "action_items" and it.get("deadline_text"):
                q, _, r = best_substring(it["deadline_text"], [{"id": "x", "text": " ".join(e["quote"] for e in it["evidence"])}])
                if r > 0.6:
                    it["deadline_text"] = q
    rev = []
    for r in m.get("reviewed_not_actionable", []):
        q, sid, _ = best_substring(r["quote"] if isinstance(r, dict) else r, segs)
        rev.append({"seg": sid, "quote": q, "reason": r.get("reason", "") if isinstance(r, dict) else ""})
    m["reviewed_not_actionable"] = rev
    return m, (sum(ratios) / len(ratios) if ratios else 0.0)


def finish_review(out: Path, m: dict[str, Any]) -> None:
    """Mô phỏng bước Claude xem lại các đoạn bị cảnh báo bỏ sót: chỉ chấp nhận đoạn là phần tách ra
    của câu đã được trích dẫn (Whisper cắt câu khác transcript gốc)."""
    from validate import validate
    rep = validate(out)
    if rep["ok"] or not rep["review_candidates"]:
        return
    segs = read_json(out / "work" / "transcript.json")["segments"]
    ids = [s["id"] for s in segs]
    cited = {e.get("seg") for sec in ("decisions", "action_items", "open_issues") for it in m[sec] for e in it["evidence"]}
    for c in rep["review_candidates"]:
        i = ids.index(c["seg"])
        neighbours = {ids[j] for j in (i - 1, i + 1) if 0 <= j < len(ids)}
        if neighbours & cited:
            m["reviewed_not_actionable"].append({"seg": c["seg"], "quote": c["text"][:60],
                                                 "reason": "Phần tách ra của câu liền kề đã được trích dẫn"})
    write_json(out / "minutes.json", m)


def apply_context_fixes(out: Path, m: dict[str, Any], gt: list[dict[str, Any]]) -> int:
    """Mô phỏng bước Claude đọc ngữ cảnh để sửa người nói cho các đoạn dẫn chứng bị tách giọng sai.
    Trong tự kiểm thử, 'ngữ cảnh' được thay bằng đáp án chuẩn của dữ liệu mẫu; khi dùng thật, Claude
    phải nêu căn cứ cụ thể (ví dụ: trả lời ngay sau khi được gọi tên)."""
    resolved_segs = {s["id"]: s for s in read_json(out / "work" / "transcript.json")["segments"]}
    rep = read_json(out / "work" / "validation.json")
    bad_ids = {e.split(":")[0].replace("Công việc ", "") for e in rep["errors"] if "không cho thấy" in e}
    ov = m.setdefault("speaker_overrides", {})
    n = 0
    for a in m["action_items"]:
        if a["id"] not in bad_ids:
            continue
        for ev in a["evidence"]:
            s = resolved_segs.get(ev.get("seg"))
            if not s:
                continue
            # người nói thật của CÂU trích dẫn (một đoạn phiên âm có thể gộp lời 2 người)
            g = max(gt, key=lambda g: SequenceMatcher(None, norm_text(ev["quote"]), norm_text(g["text"])).ratio())
            if m["speaker_map"].get(s["speaker"], {}).get("name") != g["speaker"]:
                ev["speaker"] = g["speaker"]
                ev["speaker_basis"] = "Tự kiểm thử: đối chiếu đáp án chuẩn (thay cho bước Claude đọc ngữ cảnh)"
                n += 1
    write_json(out / "minutes.json", m)
    return n


# ---------------------------------------------------------------- các kịch bản
def scenario(name: str, inputs: list[Path], base: Path, samples: dict, gt: list, skip_render_open=True) -> dict[str, Any]:
    from prepare import prepare
    from render import render
    from validate import validate
    out = base / name
    shutil.rmtree(out, ignore_errors=True)
    g = f"Kịch bản {name}"
    t0 = time.time()
    prepare([str(p) for p in inputs], out, str(samples["participants"]), "Họp kế hoạch ra mắt khóa học AI Agent",
            "2026-10-12", None, None)
    tpl = read_json(out / "minutes.template.json")
    base_min = read_json(SKILL_DIR / "samples" / "minutes-mau.json")
    info: dict[str, Any] = {"out": out}
    if tpl["speaker_map"] and all(k.startswith("SPK") or k == UNKNOWN for k in tpl["speaker_map"]):
        smap, acc = speaker_map_from_truth(out, gt)
        m, quote_sim = reanchor(base_min, out)
        m["speaker_map"] = smap
        info["diarization_accuracy"] = acc
        info["quote_similarity"] = quote_sim
        check(g, "Nhận diện người nói theo giọng", acc >= 0.75, f"đúng {acc:.0%} thời lượng nói ({len(smap)} nhãn)")
        check(g, "Phiên âm khớp kịch bản gốc", quote_sim >= 0.8, f"độ giống trung bình {quote_sim:.0%}")
    else:
        m = base_min
        m["speaker_map"] = tpl["speaker_map"]
        check(g, "Nhận diện người nói từ nhãn transcript",
              all(v["name"] != UNKNOWN for v in tpl["speaker_map"].values()), f"{len(tpl['speaker_map'])} người")
    write_json(out / "minutes.json", m)
    finish_review(out, m)
    rep = validate(out)
    owner_errs = [e for e in rep["errors"] if "không cho thấy" in e]
    if owner_errs and "diarization_accuracy" in info:
        check(g, "Kiểm chứng chặn việc gán người khi giọng bị nhận nhầm", True, f"{len(owner_errs)} công việc bị chặn")
        fixed = apply_context_fixes(out, m, gt)
        check(g, "Sửa người nói theo ngữ cảnh (speaker_overrides)", fixed > 0, f"{fixed} đoạn được sửa, có ghi căn cứ")
        rep = validate(out)
    st = rep["stats"]
    check(g, "Kiểm chứng dẫn chứng & người phụ trách", rep["ok"], "; ".join(rep["errors"][:3]))
    info["stats"] = st
    if not rep["ok"]:
        return info
    check(g, "Không bỏ sót quyết định/công việc", st["coverage_keywords_unhandled"] == 0 and st["decisions"] == 3 and st["action_items"] == 7,
          f"{st['decisions']} quyết định, {st['action_items']} công việc")
    res = render(out, open_after=False)
    pdf = out / "bien-ban.pdf"
    from render import pdf_pages
    from mmlib import load_config
    maxp = int(load_config().get("max_pages", 4))
    npg = pdf_pages(pdf) if pdf.exists() else 0
    check(g, f"Xuất PDF gọn, tối đa {maxp} trang", 1 <= npg <= maxp, f"{npg} trang")
    html_txt = (out / "bien-ban.html").read_text(encoding="utf-8")
    n_img = html_txt.count('class="row has-img"')
    check(g, "Hình dẫn chứng nằm trong phần Diễn biến", n_img > 0 and "Khoảnh khắc quan trọng" not in html_txt,
          f"{n_img} mục có hình")
    check(g, "Xuất HTML", (out / "bien-ban.html").stat().st_size > 20000)
    from openpyxl import load_workbook
    wb = load_workbook(out / "action-items.xlsx")
    n_rows = sum(1 for r in wb["Action Items"].iter_rows(min_row=5, values_only=True) if r[1])
    check(g, "Excel Action Items", n_rows == 7 and "Theo người phụ trách" in wb.sheetnames, f"{n_rows} dòng, {len(wb.sheetnames)} sheet")
    shots = [e for e in res["images"]["moments"] if e["kind"] == "screenshot"]
    ills = [e for e in res["images"]["moments"] if e["kind"] == "illustration"]
    if "video" in name:
        from PIL import Image, ImageStat
        bright = all(sum(ImageStat.Stat(Image.open(e["file"]).convert("L")).mean) > 14 for e in shots)
        check(g, "Chụp ảnh từ video tại mốc quan trọng", len(shots) == 7 and bright, f"{len(shots)} ảnh, không có khung đen")
    else:
        labelled = all("không phải ảnh chụp" in e["file"].read_text(encoding="utf-8") for e in ills)
        check(g, "Tạo hình minh hoạ khi không có video", len(ills) == 7 and labelled, f"{len(ills)} hình, đều ghi rõ là minh hoạ")
    if "audio" in name:
        check(g, "Đồ hoạ dạng sóng từ file âm thanh", (out / "hinh-anh" / "dang-song-cuoc-hop.svg").exists())
    plan = read_json(out / "email" / "email_plan.json")["emails"]
    long_ = next((e for e in plan if e["name"] == "Phạm Đức Long"), None)
    payload = read_json(out / long_["payload"]) if long_ else {}
    check(g, "Email riêng cho từng người phụ trách", len(plan) == 3 and all((out / e["eml"]).exists() for e in plan),
          ", ".join(e["name"] for e in plan))
    check(g, "Thiếu email → đánh dấu, không điền người nhận", long_ and long_["status"] == "thiếu email" and payload.get("to") == [],
          "Phạm Đức Long: " + (long_["status"] if long_ else "?"))
    check(g, "Chưa gửi email nào (chờ phê duyệt)", all(e["status"] in ("chờ phê duyệt", "thiếu email") for e in plan))
    info["seconds"] = round(time.time() - t0)
    return info


def negative_tests(base: Path, samples: dict) -> None:
    from prepare import load_transcript, prepare
    from render import render
    from validate import validate
    g = "Xử lý lỗi"
    src = base / "transcript"
    neg = base / "_loi"
    shutil.rmtree(neg, ignore_errors=True)
    shutil.copytree(src, neg)
    good = read_json(src / "minutes.json")

    def with_change(fn):
        m = copy.deepcopy(good)
        fn(m)
        write_json(neg / "minutes.json", m)
        return validate(neg)

    r = with_change(lambda m: m["decisions"].append({"id": "QĐ9", "text": "Tăng ngân sách quảng cáo lên 200 triệu",
                                                      "evidence": [{"quote": "chúng ta tăng ngân sách quảng cáo lên hai trăm triệu"}]}))
    check(g, "Phát hiện dẫn chứng bịa", not r["ok"] and any("KHÔNG tìm thấy" in e for e in r["errors"]))
    r = with_change(lambda m: m["action_items"][1].update(owner="Hoàng Văn Nam"))
    check(g, "Người phụ trách không có trong danh sách", not r["ok"] and any("không có trong danh sách" in e for e in r["errors"]))
    r = with_change(lambda m: m["action_items"][1].update(owner="Lê Thu Hà"))
    check(g, "Gán sai người phụ trách (dẫn chứng không ủng hộ)", not r["ok"] and any("không cho thấy" in e for e in r["errors"]))
    r = with_change(lambda m: m["action_items"][5].update(deadline="2026-10-22", deadline_text="thứ Năm tuần sau"))
    check(g, "Thời hạn tự suy đoán", not r["ok"] and any("không xuất hiện trong dẫn chứng" in e for e in r["errors"]))
    r = with_change(lambda m: m["action_items"].pop(1))
    check(g, "Phát hiện bỏ sót công việc", not r["ok"] and any(c["keyword"] for c in r["review_candidates"]),
          f"{len(r['review_candidates'])} đoạn bị cảnh báo")
    expect_error(g, "Chặn xuất bản khi kiểm chứng chưa đạt", lambda: render(neg, open_after=False), "chưa")
    expect_error(g, "File không tồn tại", lambda: prepare([str(base / "khong-co.mp3")], neg, None, "", "", None, None), "không tìm thấy")
    bad = base / "sai-dinh-dang.xyz"
    bad.write_text("abc", encoding="utf-8")
    expect_error(g, "Định dạng không hỗ trợ", lambda: prepare([str(bad)], neg, None, "", "", None, None), "chưa hỗ trợ")
    broken = base / "hong.mp3"
    broken.write_bytes(b"\x00\x01khong-phai-am-thanh" * 200)
    expect_error(g, "File âm thanh hỏng", lambda: prepare([str(broken)], neg, None, "", "", None, None), "hỏng")
    empty = base / "rong.txt"
    empty.write_text("", encoding="utf-8")
    expect_error(g, "File rỗng", lambda: prepare([str(empty)], neg, None, "", "", None, None), "rỗng")
    bad.unlink()
    broken.unlink()
    empty.unlink()
    # định dạng transcript khác nhau
    tmp = base / "_fmt"
    tmp.mkdir(exist_ok=True)
    (tmp / "a.vtt").write_text("WEBVTT\n\n00:00:01.000 --> 00:00:04.000\n<v Lê Thu Hà>Em sẽ gửi báo cáo.</v>\n\n"
                               "00:00:05.000 --> 00:00:07.000\n<v Trần Quốc Bảo>Dạ vâng.</v>\n", encoding="utf-8")
    (tmp / "b.srt").write_text("1\n00:00:01,000 --> 00:00:03,000\nPhạm Đức Long: Mình đồng ý.\n", encoding="utf-8")
    (tmp / "c.txt").write_text("Nguyễn Minh Anh (00:01:02): Chốt phương án A.\nTrần Quốc Bảo (00:01:10): Em nhận.\n", encoding="utf-8")
    (tmp / "d.txt").write_text("Nguyễn Minh Anh\n00:00:05\nChào mọi người.\nLê Thu Hà\n00:00:09\nEm báo cáo nhanh.\n", encoding="utf-8")
    ok = True
    for f, exp in [("a.vtt", ("Lê Thu Hà", 1.0)), ("b.srt", ("Phạm Đức Long", 1.0)), ("c.txt", ("Nguyễn Minh Anh", 62.0)),
                   ("d.txt", ("Nguyễn Minh Anh", 5.0))]:
        s = load_transcript(tmp / f)
        ok &= bool(s) and s[0]["speaker"] == exp[0] and abs((s[0]["start"] or 0) - exp[1]) < 0.01
    check(g, "Đọc nhiều định dạng transcript (VTT Teams, SRT Zoom, TXT, Google Meet)", ok)
    shutil.rmtree(tmp)
    # lọc câu ảo giác của Whisper
    import re as _re
    from mmlib import load_config
    pats = load_config().get("hallucination_patterns", [])
    t = " Hãy subscribe cho kênh Ghiền Mì Gõ Để không bỏ lỡ những video hấp dẫn thì mô hình sẽ không thể sử dụng"
    for p in pats:
        t = _re.sub(p, "", t, flags=_re.IGNORECASE)
    check(g, "Lọc câu ảo giác của Whisper (\"subscribe cho kênh…\")", t.strip() == "thì mô hình sẽ không thể sử dụng", t.strip()[:60])
    # bổ sung email
    from mm import main as _main  # noqa: F401  (đảm bảo CLI import được)
    from mmlib import EMAIL_RE
    check(g, "Từ chối email sai định dạng", not EMAIL_RE.fullmatch("long@@abc"))
    ps = read_json(src / "work" / "participants.json")
    for p in ps:
        if p["name"] == "Phạm Đức Long":
            p["email"], p["email_valid"] = "long.pham@example.com", True
    write_json(neg / "work" / "participants.json", ps)
    write_json(neg / "minutes.json", good)
    validate(neg)
    render(neg, open_after=False, only_emails=True)
    plan = read_json(neg / "email" / "email_plan.json")["emails"]
    long_ = next(e for e in plan if e["name"] == "Phạm Đức Long")
    check(g, "Bổ sung email → email được điền người nhận, vẫn chờ phê duyệt",
          long_["email"] == "long.pham@example.com" and long_["status"] == "chờ phê duyệt")
    shutil.rmtree(neg, ignore_errors=True)


def write_report(base: Path, infos: dict[str, Any], seconds: float) -> Path:
    ok = sum(r["ok"] for r in RESULTS)
    lines = ["# Báo cáo tự kiểm thử Agent Biên bản họp (theo PEAS)", "",
             f"Thời điểm: {time.strftime('%Y-%m-%d %H:%M')} · Tổng thời gian: {seconds:.0f} giây · "
             f"**Kết quả: {ok}/{len(RESULTS)} ca đạt**", ""]
    lines += ["## P — Kết quả (Performance)", "", "| Tiêu chí | Kịch bản | Kết quả |", "|---|---|---|"]
    for name, inf in infos.items():
        st = inf.get("stats", {})
        if st:
            lines.append(f"| Dẫn chứng khớp transcript | {name} | {st['evidence_exact'] + st['evidence_near']}/{st['evidence_total']} "
                         f"(không tìm thấy: {st['evidence_failed']}) |")
            lines.append(f"| Quyết định / công việc ghi nhận (đáp án: 3 / 7) | {name} | {st['decisions']} / {st['action_items']} |")
            lines.append(f"| Công việc có người phụ trách (đáp án: 6, 1 chưa xác định) | {name} | {st['actions_with_owner']} |")
        if "diarization_accuracy" in inf:
            lines.append(f"| Nhận diện người nói theo giọng | {name} | {inf['diarization_accuracy']:.0%} thời lượng đúng |")
            lines.append(f"| Độ giống của phiên âm so với lời thoại gốc | {name} | {inf['quote_similarity']:.0%} |")
    lines += ["", "## E — Môi trường (Environment)", "",
              "- Thư mục máy tính: đọc file mẫu trong `samples/`, ghi kết quả vào thư mục kiểm thử.",
              "- Gmail / Google Drive: kiểm tra ở bước gửi thật (cần người dùng kết nối & phê duyệt); tự kiểm thử KHÔNG gửi email.", "",
              "## A — Hành động (Actions) & S — Đầu vào (Sensors)", "", "| Nhóm | Ca kiểm thử | Đạt | Chi tiết |", "|---|---|---|---|"]
    for r in RESULTS:
        lines.append(f"| {r['group']} | {r['name']} | {'✅' if r['ok'] else '❌'} | {r['detail'].replace('|', '/')} |")
    lines += ["", "## Kết quả mẫu để xem", ""]
    for name, inf in infos.items():
        lines.append(f"- {name}: `{inf['out']}` (bien-ban.html, bien-ban.pdf, action-items.xlsx, email/xem-truoc-email.html)")
    p = base / "bao-cao-tu-kiem-thu.md"
    p.write_text("\n".join(lines) + "\n", encoding="utf-8")
    write_json(base / "bao-cao-tu-kiem-thu.json", RESULTS)
    return p


def run_selftest(skip_media: bool = False, base: Path | None = None) -> int:
    from make_samples import make_samples
    t0 = time.time()
    base = (base or Path.cwd() / "bien-ban-tu-kiem-thu").resolve()
    base.mkdir(parents=True, exist_ok=True)
    log(f"🧪 TỰ KIỂM THỬ — kết quả lưu tại {base}")
    samples = make_samples()
    gt = read_json(samples["ground_truth"])["segments"]
    infos = {}
    log("\n▶ Kịch bản 1: transcript (.txt) + danh sách người tham gia")
    infos["transcript"] = scenario("transcript", [samples["transcript"]], base, samples, gt)
    if not skip_media:
        log("\n▶ Kịch bản 2: chỉ có video (.mp4) — phiên âm + nhận diện giọng + chụp ảnh")
        infos["video"] = scenario("video", [samples["mp4"]], base, samples, gt)
        log("\n▶ Kịch bản 3: chỉ có ghi âm (.mp3) — phiên âm + hình minh hoạ")
        infos["audio"] = scenario("audio", [samples["mp3"]], base, samples, gt)
    log("\n▶ Kịch bản 4: các tình huống lỗi")
    negative_tests(base, samples)
    rep = write_report(base, infos, time.time() - t0)
    ok = sum(r["ok"] for r in RESULTS)
    log(f"\n{'✅' if ok == len(RESULTS) else '⚠'} {ok}/{len(RESULTS)} ca đạt — báo cáo: {rep}")
    return 0 if ok == len(RESULTS) else 1
