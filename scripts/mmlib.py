"""Tiện ích dùng chung cho Agent biên bản cuộc họp (bien-ban-hop)."""
from __future__ import annotations

import csv
import json
import re
import shutil
import sys
import unicodedata
import zipfile
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

SKILL_DIR = Path(__file__).resolve().parent.parent
ASSETS = SKILL_DIR / "assets"
MODELS = SKILL_DIR / "models"

AUDIO_EXT = {".mp3", ".wav", ".m4a", ".aac", ".ogg", ".flac", ".wma", ".opus"}
VIDEO_EXT = {".mp4", ".mov", ".mkv", ".webm", ".avi", ".m4v", ".wmv"}
TEXT_EXT = {".txt", ".md", ".vtt", ".srt", ".docx", ".json"}
UNKNOWN = "Chưa xác định"

# Windows console mặc định cp1252 → ép UTF-8 để in tiếng Việt.
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


class AgentError(Exception):
    """Lỗi có thông điệp thân thiện cho người dùng (tiếng Việt) và gợi ý cách xử lý."""

    def __init__(self, message: str, hint: str = ""):
        super().__init__(message)
        self.hint = hint


def log(msg: str) -> None:
    print(msg, flush=True)


def load_config() -> dict[str, Any]:
    cfg = read_json(SKILL_DIR / "config.json")
    user_cfg = SKILL_DIR / "config.local.json"
    if user_cfg.exists():
        cfg.update(read_json(user_cfg))
    return cfg


def read_json(path: Path) -> Any:
    try:
        return json.loads(Path(path).read_text(encoding="utf-8-sig"))
    except FileNotFoundError:
        raise AgentError(f"Không tìm thấy file: {path}")
    except json.JSONDecodeError as e:
        raise AgentError(f"File JSON bị lỗi cú pháp: {path} (dòng {e.lineno}, cột {e.colno}): {e.msg}",
                         "Mở file, sửa dấu phẩy/ngoặc bị thiếu rồi chạy lại.")


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


# ---------------------------------------------------------------- thời gian
def ts_to_sec(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(",", ".")
    m = re.fullmatch(r"(?:(\d+):)?(\d{1,2}):(\d{1,2}(?:\.\d+)?)", s)
    if not m:
        return None
    h = int(m.group(1) or 0)
    return h * 3600 + int(m.group(2)) * 60 + float(m.group(3))


def sec_to_ts(sec: float | None) -> str:
    if sec is None:
        return ""
    sec = max(0, int(round(sec)))
    return f"{sec // 3600:02d}:{sec % 3600 // 60:02d}:{sec % 60:02d}"


# ---------------------------------------------------------------- chuẩn hoá chữ
_WEEKDAYS = {"hai": "2", "ba": "3", "tư": "4", "tu": "4", "năm": "5", "sáu": "6", "bảy": "7"}


def norm_text(s: str) -> str:
    """Chuẩn hoá để so khớp dẫn chứng: NFC, chữ thường, bỏ dấu câu, gộp khoảng trắng."""
    s = unicodedata.normalize("NFC", s or "").lower()
    s = re.sub(r"thứ\s+(hai|ba|tư|năm|sáu|bảy)\b", lambda m: "thứ " + _WEEKDAYS[m.group(1)], s)
    s = re.sub(r"(\d)[.,](?=\d{3}\b)", r"\1", s)  # 1.490.000 -> 1490000
    s = re.sub(r"[^\w\s]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def strip_accents(s: str) -> str:
    s = unicodedata.normalize("NFD", s or "").replace("đ", "d").replace("Đ", "D")
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def slugify(s: str, maxlen: int = 50) -> str:
    s = re.sub(r"[^a-zA-Z0-9]+", "-", strip_accents(s).lower()).strip("-")
    return (s[:maxlen].strip("-")) or "khong-ten"


def partial_ratio(needle: str, hay: str) -> float:
    """Tỉ lệ khớp tốt nhất của `needle` bên trong `hay` (0..1), đã chuẩn hoá."""
    n, h = norm_text(needle), norm_text(hay)
    if not n:
        return 0.0
    if n in h:
        return 1.0
    nw, hw = n.split(), h.split()
    L = len(nw)
    if L == 0 or not hw:
        return 0.0
    best = 0.0
    nset = set(nw)
    for size in {max(1, L - 2), L, L + 2}:
        for i in range(0, max(1, len(hw) - size + 1)):
            win = hw[i:i + size]
            if len(nset.intersection(win)) < 0.4 * L:
                continue
            r = SequenceMatcher(None, n, " ".join(win), autojunk=False).ratio()
            if r > best:
                best = r
                if best > 0.995:
                    return best
    return best


# ---------------------------------------------------------------- người tham gia
def load_participants(path: Path | None) -> list[dict[str, Any]]:
    """Đọc danh sách người tham gia từ CSV / XLSX / JSON / TXT.

    Cột nhận diện được: name|họ tên|tên, email, role|vai trò|chức vụ, aliases|biệt danh (phân cách bằng ;)
    """
    if not path:
        return []
    path = Path(path)
    if not path.exists():
        raise AgentError(f"Không tìm thấy file danh sách người tham gia: {path}")
    ext = path.suffix.lower()
    rows: list[dict[str, Any]] = []
    if ext == ".json":
        data = read_json(path)
        rows = data.get("participants", data) if isinstance(data, dict) else data
    elif ext in {".csv", ".tsv"}:
        text = path.read_text(encoding="utf-8-sig")
        dialect = csv.Sniffer().sniff(text.splitlines()[0], delimiters=",;\t") if text.strip() else csv.excel
        rows = list(csv.DictReader(text.splitlines(), dialect=dialect))
    elif ext in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(it)]
        rows = [dict(zip(header, r)) for r in it if any(r)]
    elif ext == ".txt":
        for line in path.read_text(encoding="utf-8-sig").splitlines():
            if not line.strip():
                continue
            parts = [p.strip() for p in re.split(r"\s*[|,;\t]\s*", line.strip())]
            row = {"name": parts[0]}
            for p in parts[1:]:
                if "@" in p:
                    row["email"] = p
                elif p:
                    row.setdefault("role", p)
            rows.append(row)
    else:
        raise AgentError(f"Định dạng danh sách người tham gia chưa hỗ trợ: {ext}",
                         "Dùng CSV, XLSX, JSON hoặc TXT (mỗi dòng: Họ tên, email, vai trò).")

    keymap = {"name": ["name", "họ tên", "ho ten", "tên", "ten", "họ và tên", "full name"],
              "email": ["email", "e-mail", "mail", "địa chỉ email"],
              "role": ["role", "vai trò", "vai tro", "chức vụ", "chuc vu", "bộ phận", "title"],
              "aliases": ["aliases", "alias", "biệt danh", "tên gọi khác", "ten goi khac"]}
    out = []
    for r in rows:
        low = {str(k or "").strip().lower(): v for k, v in dict(r).items()}
        p: dict[str, Any] = {}
        for key, names in keymap.items():
            for n in names:
                if low.get(n) not in (None, ""):
                    p[key] = low[n]
                    break
        name = str(p.get("name", "")).strip()
        if not name:
            continue
        aliases = p.get("aliases") or []
        if isinstance(aliases, str):
            aliases = [a.strip() for a in re.split(r"[;,]", aliases) if a.strip()]
        email = str(p.get("email") or "").strip()
        out.append({"name": name, "email": email, "role": str(p.get("role") or "").strip(),
                    "aliases": aliases, "email_valid": bool(EMAIL_RE.fullmatch(email)) if email else False})
    return out


EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")


def participant_lookup(participants: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    look = {}
    for p in participants:
        keys = [p["name"], *p.get("aliases", [])]
        given = p["name"].split()[-1] if p["name"].split() else ""
        keys.append(given)
        for k in keys:
            if k:
                look.setdefault(norm_text(k), p)
    return look


# ---------------------------------------------------------------- công cụ ngoài
def find_ffmpeg(name: str = "ffmpeg") -> str | None:
    exe = shutil.which(name)
    if exe:
        return exe
    base = Path.home() / "AppData/Local/Microsoft/WinGet/Packages"
    if base.exists():
        for p in base.glob(f"*FFmpeg*/**/bin/{name}.exe"):
            return str(p)
    return None


def find_browser() -> str | None:
    cands = [
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
        str(Path.home() / r"AppData\Local\Google\Chrome\Application\chrome.exe"),
        "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
        "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
    ]
    for c in cands:
        if Path(c).exists():
            return c
    for n in ("msedge", "google-chrome", "chromium", "chromium-browser", "chrome"):
        if shutil.which(n):
            return shutil.which(n)
    return None


def read_docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    xml = re.sub(r"</w:p>", "\n", xml)
    xml = re.sub(r"<w:tab/>", "\t", xml)
    xml = re.sub(r"<w:br/>", "\n", xml)
    text = re.sub(r"<[^>]+>", "", xml)
    import html as _h
    return _h.unescape(text)
