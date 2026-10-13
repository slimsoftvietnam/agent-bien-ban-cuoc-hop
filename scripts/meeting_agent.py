#!/usr/bin/env python3
"""Conservative meeting-minutes pipeline for the meeting-minutes-agent skill."""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import os
import re
import shutil
import subprocess
import sys
import time
import unicodedata
import zipfile
from collections import defaultdict
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path
from typing import Any


SKILL_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = SKILL_ROOT / "config" / "default.json"
MEDIA_EXTS = {".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg", ".mp4", ".mov", ".mkv", ".webm"}
VIDEO_EXTS = {".mp4", ".mov", ".mkv", ".webm"}
LOCAL_RUNTIME = SKILL_ROOT / ".runtime"

if LOCAL_RUNTIME.is_dir() and str(LOCAL_RUNTIME) not in sys.path:
    sys.path.insert(0, str(LOCAL_RUNTIME))

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")


def fail(message: str, code: int = 2) -> None:
    print(f"ERROR: {message}", file=sys.stderr)
    raise SystemExit(code)


def read_json(path: Path) -> dict[str, Any]:
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except Exception as exc:
        fail(f"Không đọc được JSON {path}: {exc}")


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalize_key(value: str) -> str:
    decomposed = unicodedata.normalize("NFD", value.strip().casefold())
    return "".join(ch for ch in decomposed if unicodedata.category(ch) != "Mn")


def safe_name(value: str) -> str:
    clean = normalize_key(value)
    clean = re.sub(r"[^a-z0-9]+", "-", clean).strip("-")
    return clean or "chua-xac-dinh"


def timestamp_to_seconds(value: str | None) -> float | None:
    if not value:
        return None
    parts = [int(part) for part in value.split(":")]
    if len(parts) == 2:
        return float(parts[0] * 60 + parts[1])
    if len(parts) == 3:
        return float(parts[0] * 3600 + parts[1] * 60 + parts[2])
    return None


def seconds_to_timestamp(value: float) -> str:
    total = max(0, int(round(value)))
    return f"{total // 3600:02d}:{(total % 3600) // 60:02d}:{total % 60:02d}"


def load_participants(path: Path) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]]]:
    payload = read_json(path)
    participants = payload.get("participants")
    if not isinstance(participants, list):
        fail("participants.json phải có khóa 'participants' là một danh sách.")
    alias_map: dict[str, dict[str, Any]] = {}
    for raw in participants:
        if not isinstance(raw, dict) or not str(raw.get("name", "")).strip():
            fail("Mỗi người tham gia phải có trường name.")
        item = {
            "name": str(raw["name"]).strip(),
            "email": str(raw.get("email", "")).strip(),
            "aliases": [str(v).strip() for v in raw.get("aliases", []) if str(v).strip()],
        }
        raw.clear()
        raw.update(item)
        for alias in [item["name"], *item["aliases"]]:
            key = normalize_key(alias)
            if key in alias_map and alias_map[key]["name"] != item["name"]:
                fail(f"Alias trùng giữa nhiều người: {alias}")
            alias_map[key] = item
    return participants, alias_map


def decode_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "utf-16", "cp1258", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    fail(f"Không xác định được encoding của {path}")
    return ""


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in {".txt", ".md"}:
        return decode_text(path)
    if ext == ".docx":
        try:
            from docx import Document
        except ImportError:
            fail("Thiếu python-docx. Hãy dùng Python bundled của Codex hoặc cài python-docx.")
        doc = Document(str(path))
        chunks = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(chunks)
        if not text.strip():
            fail("DOCX không có nội dung chữ có thể trích xuất.")
        return text
    if ext == ".pdf":
        try:
            import pdfplumber
        except ImportError:
            fail("Thiếu pdfplumber. Hãy dùng Python bundled của Codex hoặc cài pdfplumber.")
        pages: list[str] = []
        with pdfplumber.open(str(path)) as pdf:
            for index, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                if text.strip():
                    pages.append(f"[Trang {index}]\n{text}")
        if not pages:
            fail("PDF không có lớp chữ. Hãy OCR tài liệu trước khi chạy Agent.")
        return "\n".join(pages)
    fail(f"Định dạng transcript chưa được hỗ trợ: {ext}")
    return ""


def transcribe_media(path: Path, output_dir: Path, model_name: str) -> Path:
    try:
        from faster_whisper import WhisperModel
    except ImportError:
        fail(
            "Đầu vào là audio/video nhưng chưa có transcript và faster-whisper chưa được cài. "
            "Hãy truyền --transcript hoặc cài faster-whisper."
        )
    print(f"Đang phiên âm cục bộ bằng faster-whisper model={model_name}...")
    model = WhisperModel(model_name, device="cpu", compute_type="int8")
    segments, _ = model.transcribe(str(path), vad_filter=True)
    lines = []
    for segment in segments:
        stamp = seconds_to_timestamp(float(segment.start))
        lines.append(f"[{stamp}] Chưa xác định: {segment.text.strip()}")
    if not lines:
        fail("Whisper không tạo được segment transcript nào.")
    generated = output_dir / "analysis" / "generated-transcript.txt"
    generated.parent.mkdir(parents=True, exist_ok=True)
    generated.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return generated


LINE_RE = re.compile(
    r"^\s*(?:\[(?P<bracket>\d{1,2}:\d{2}(?::\d{2})?)\]|(?P<plain>\d{1,2}:\d{2}(?::\d{2})?))?\s*"
    r"(?P<speaker>[^:\n]{1,80})\s*:\s*(?P<text>.+?)\s*$"
)


def parse_segments(text: str, alias_map: dict[str, dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    segments: list[dict[str, Any]] = []
    reviews: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    for line_no, raw_line in enumerate(text.splitlines(), start=1):
        line = raw_line.strip()
        if not line or re.fullmatch(r"\[Trang \d+\]", line):
            continue
        match = LINE_RE.match(line)
        if match:
            source_speaker = match.group("speaker").strip()
            participant = alias_map.get(normalize_key(source_speaker))
            stamp = match.group("bracket") or match.group("plain")
            current = {
                "line": line_no,
                "timestamp": stamp,
                "timestamp_seconds": timestamp_to_seconds(stamp),
                "source_speaker": source_speaker,
                "speaker": participant["name"] if participant else "Chưa xác định",
                "speaker_email": participant["email"] if participant else "",
                "text": match.group("text").strip(),
                "evidence_ref": stamp or f"Dòng {line_no}",
                "speaker_confirmed": bool(participant),
            }
            segments.append(current)
            if not participant:
                reviews.append(
                    {
                        "type": "speaker",
                        "evidence_ref": current["evidence_ref"],
                        "message": f"Không khớp nhãn người nói '{source_speaker}' với danh sách người tham gia.",
                    }
                )
        elif current:
            current["text"] += " " + line
        else:
            current = {
                "line": line_no,
                "timestamp": None,
                "timestamp_seconds": None,
                "source_speaker": "Chưa xác định",
                "speaker": "Chưa xác định",
                "speaker_email": "",
                "text": line,
                "evidence_ref": f"Dòng {line_no}",
                "speaker_confirmed": False,
            }
            segments.append(current)
            reviews.append(
                {
                    "type": "format",
                    "evidence_ref": current["evidence_ref"],
                    "message": "Dòng không có cấu trúc người nói; không suy đoán danh tính.",
                }
            )
    if not segments:
        fail("Không tìm thấy nội dung transcript có thể xử lý.")
    return segments, reviews


def find_deadline(text: str, patterns: list[str]) -> tuple[str, str | None]:
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if not match:
            continue
        raw = match.group(0).strip()
        iso_value: str | None = None
        for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
            try:
                iso_value = datetime.strptime(raw, fmt).date().isoformat()
                break
            except ValueError:
                continue
        return raw, iso_value
    return "Chưa nêu", None


def match_owners(text: str, segment: dict[str, Any], participants: list[dict[str, Any]]) -> list[dict[str, Any]]:
    normalized = normalize_key(text)
    candidates: list[dict[str, Any]] = []
    if re.search(r"\b(toi|em|anh|chi|minh)\s+se\b", normalized) and segment["speaker_confirmed"]:
        candidates.append({"name": segment["speaker"], "email": segment["speaker_email"]})
    for person in participants:
        aliases = sorted([person["name"], *person.get("aliases", [])], key=len, reverse=True)
        for alias in aliases:
            key = re.escape(normalize_key(alias))
            owner_pattern = rf"(?:\b{key}\b.{{0,35}}\b(?:phu trach|chiu trach nhiem|se|cam ket)\b|\b(?:giao|phan cong)\b.{{0,25}}\b{key}\b)"
            if re.search(owner_pattern, normalized):
                if not any(item["name"] == person["name"] for item in candidates):
                    candidates.append({"name": person["name"], "email": person["email"]})
                break
    return candidates


def clean_task(text: str, owners: list[dict[str, Any]], deadline_text: str) -> str:
    task = text.strip()
    task = re.sub(r"^(?:tôi|em|anh|chị|mình)\s+sẽ\s+", "", task, flags=re.IGNORECASE)
    task = re.sub(r"^.{1,80}?\s+(?:phụ trách|chịu trách nhiệm|sẽ|cam kết)\s+", "", task, count=1, flags=re.IGNORECASE)
    for owner in owners:
        aliases = [owner["name"]]
        for alias in aliases:
            task = re.sub(
                rf"^{re.escape(alias)}\s+(?:sẽ\s+|phụ trách\s+|chịu trách nhiệm\s+|cam kết\s+)",
                "",
                task,
                flags=re.IGNORECASE,
            )
    if deadline_text != "Chưa nêu":
        task = re.sub(
            rf"(?:,?\s*(?:(?:trước\s+ngày)|hạn|ngày)?\s*)?{re.escape(deadline_text)}[.!]?\s*$",
            "",
            task,
            flags=re.IGNORECASE,
        )
    return task.strip(" .,-") or text.strip()


def analyze(
    segments: list[dict[str, Any]],
    participants: list[dict[str, Any]],
    config: dict[str, Any],
    reviews: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    decisions: list[dict[str, Any]] = []
    actions: list[dict[str, Any]] = []
    max_quote = int(config.get("max_evidence_quote_chars", 320))
    for segment in segments:
        lowered = segment["text"].casefold()
        if any(keyword.casefold() in lowered for keyword in config["decision_keywords"]):
            decision = {
                "id": f"D{len(decisions) + 1:02d}",
                "decision": segment["text"],
                "speaker": segment["speaker"],
                "evidence_quote": segment["text"][:max_quote],
                "evidence_ref": segment["evidence_ref"],
                "timestamp_seconds": segment["timestamp_seconds"],
                "review_status": "Đã xác định" if segment["speaker_confirmed"] else "Cần rà soát",
            }
            decisions.append(decision)
        if not any(keyword.casefold() in lowered for keyword in config["action_keywords"]):
            continue
        owners = match_owners(segment["text"], segment, participants)
        deadline_text, deadline_iso = find_deadline(segment["text"], config["deadline_patterns"])
        owner = owners[0] if len(owners) == 1 else {"name": "Chưa xác định", "email": ""}
        issues: list[str] = []
        if len(owners) == 0:
            issues.append("Không có người phụ trách được nêu rõ")
        elif len(owners) > 1:
            issues.append("Có nhiều người có thể là người phụ trách")
        if deadline_text == "Chưa nêu":
            issues.append("Không có deadline")
        if owner["name"] != "Chưa xác định" and not owner["email"]:
            issues.append("Người phụ trách chưa có email")
        item = {
            "id": f"A{len(actions) + 1:02d}",
            "task": clean_task(segment["text"], owners, deadline_text),
            "owner": owner["name"],
            "owner_email": owner["email"],
            "deadline_text": deadline_text,
            "deadline_iso": deadline_iso,
            "evidence_quote": segment["text"][:max_quote],
            "evidence_ref": segment["evidence_ref"],
            "timestamp_seconds": segment["timestamp_seconds"],
            "review_status": "Cần rà soát" if issues else "Đã xác định",
            "review_reasons": issues,
        }
        actions.append(item)
        for issue in issues:
            reviews.append({"type": "action_item", "item_id": item["id"], "evidence_ref": item["evidence_ref"], "message": issue})
    return decisions, actions


def capture_screenshots(video: Path | None, decisions: list[dict[str, Any]], actions: list[dict[str, Any]], output_dir: Path, width: int) -> tuple[list[dict[str, Any]], list[str]]:
    screenshots: list[dict[str, Any]] = []
    warnings: list[str] = []
    if not video:
        return screenshots, warnings
    ffmpeg = shutil.which("ffmpeg")
    if not ffmpeg:
        warnings.append("Không tìm thấy FFmpeg; không tạo được screenshot từ video.")
        return screenshots, warnings
    shots_dir = output_dir / "screenshots"
    shots_dir.mkdir(parents=True, exist_ok=True)
    seen: set[int] = set()
    for kind, items in (("quyet-dinh", decisions), ("cong-viec", actions)):
        for item in items:
            second = item.get("timestamp_seconds")
            if second is None:
                continue
            rounded = int(round(float(second)))
            if rounded in seen:
                continue
            seen.add(rounded)
            filename = f"moc-{rounded:06d}-{kind}-{item['id']}.png"
            target = shots_dir / filename
            command = [
                ffmpeg,
                "-hide_banner",
                "-loglevel",
                "error",
                "-ss",
                str(max(0, float(second))),
                "-i",
                str(video),
                "-frames:v",
                "1",
                "-vf",
                f"scale={width}:-2",
                "-y",
                str(target),
            ]
            result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace")
            if result.returncode != 0 or not target.exists():
                warnings.append(f"Không chụp được mốc {item['evidence_ref']} từ video: {result.stderr.strip()}")
                continue
            screenshots.append(
                {
                    "path": str(target.relative_to(output_dir)).replace("\\", "/"),
                    "timestamp": seconds_to_timestamp(float(second)),
                    "timestamp_seconds": second,
                    "kind": kind,
                    "item_id": item["id"],
                    "caption": item.get("decision") or item.get("task"),
                }
            )
    return screenshots, warnings


def html_table(headers: list[str], rows: list[list[str]]) -> str:
    head = "".join(f"<th>{html.escape(value)}</th>" for value in headers)
    body = "".join("<tr>" + "".join(f"<td>{value}</td>" for value in row) + "</tr>" for row in rows)
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def build_html(data: dict[str, Any], output_dir: Path) -> Path:
    css = (SKILL_ROOT / "assets" / "report.css").read_text(encoding="utf-8")
    decisions = data["decisions"]
    actions = data["action_items"]
    participants = data["participants"]
    participant_list = "".join(f"<span class='person'>{html.escape(p['name'])}</span>" for p in participants) or "<span class='muted'>Chưa cung cấp</span>"
    duration_seconds = max((segment.get("timestamp_seconds") or 0 for segment in data["segments"]), default=0)
    duration_label = seconds_to_timestamp(duration_seconds) if duration_seconds else "Chưa xác định"

    quick_items: list[str] = []
    for item in decisions[:3]:
        quick_items.append(f"<li><strong>Quyết định:</strong> {html.escape(item['decision'])}</li>")
    for item in actions[:3]:
        quick_items.append(
            f"<li><strong>Việc cần làm:</strong> {html.escape(item['task'])} - "
            f"{html.escape(item['owner'])}, hạn {html.escape(item['deadline_text'])}</li>"
        )
    if not quick_items:
        quick_items.append("<li>Không có quyết định hoặc công việc rõ ràng trong nguồn.</li>")

    action_cards = []
    for item in actions:
        status_class = " review-needed" if item["review_status"] != "Đã xác định" else ""
        action_cards.append(
            f"<article class='task{status_class}' id='{html.escape(item['id'].lower())}' "
            f"data-owner='{html.escape(item['owner'], quote=True)}' "
            f"data-deadline='{html.escape(item.get('deadline_iso') or '9999-12-31', quote=True)}'>"
            f"<div class='task-head'><span class='task-id'>{html.escape(item['id'])}</span>"
            f"<div class='task-title'>{html.escape(item['task'])}</div>"
            f"<span class='chip'>{html.escape(item['review_status'])}</span></div>"
            f"<div class='task-meta'><strong>Phụ trách:</strong> {html.escape(item['owner'])} &nbsp; "
            f"<strong>Deadline:</strong> {html.escape(item['deadline_text'])}</div>"
            f"<div class='evidence'><strong>{html.escape(item['evidence_ref'])}</strong> - “{html.escape(item['evidence_quote'])}”</div>"
            "</article>"
        )
    if not action_cards:
        action_cards.append("<p>Không có Action Item rõ ràng trong nguồn.</p>")

    decision_cards = []
    for item in decisions:
        decision_cards.append(
            f"<article class='decision' id='{html.escape(item['id'].lower())}'>"
            f"<div class='task-head'><span class='task-id'>{html.escape(item['id'])}</span>"
            f"<div class='task-title'>{html.escape(item['decision'])}</div></div>"
            f"<div class='task-meta'><strong>Người nói:</strong> {html.escape(item['speaker'])} &nbsp; "
            f"<strong>Trạng thái:</strong> {html.escape(item['review_status'])}</div>"
            f"<div class='evidence'><strong>{html.escape(item['evidence_ref'])}</strong> - “{html.escape(item['evidence_quote'])}”</div>"
            "</article>"
        )
    if not decision_cards:
        decision_cards.append("<p>Không có quyết định rõ ràng trong nguồn.</p>")

    figures = "".join(
        f"<figure><img src='{html.escape(shot['path'])}' alt='Ảnh mốc {html.escape(shot['timestamp'])}'>"
        f"<figcaption>{html.escape(shot['timestamp'])} · {html.escape(shot['caption'])}</figcaption></figure>"
        for shot in data["screenshots"]
    )
    speakers = sorted({segment["speaker"] for segment in data["segments"]})
    speaker_buttons = "".join(
        f"<button type='button' data-speaker-filter='{html.escape(name, quote=True)}'>{html.escape(name)}</button>" for name in speakers
    )
    timeline = []
    for segment in data["segments"]:
        anchor = f"t-{int(segment['timestamp_seconds']):04d}" if segment.get("timestamp_seconds") is not None else f"line-{segment['line']}"
        timeline.append(
            f"<article class='turn' id='{anchor}' data-speaker='{html.escape(segment['speaker'], quote=True)}'>"
            f"<div class='turn-time'>{html.escape(segment['evidence_ref'])}</div>"
            f"<div class='turn-speaker'>{html.escape(segment['speaker'])}</div>"
            f"<div class='turn-text'>{html.escape(segment['text'])}</div></article>"
        )

    seen_reviews: set[str] = set()
    review_rows: list[str] = []
    for item in data["review_queue"]:
        message = str(item["message"])
        if message in seen_reviews:
            continue
        seen_reviews.add(message)
        review_rows.append(
            f"<li>{html.escape(message)} <span class='muted'>({html.escape(item.get('evidence_ref', ''))})</span></li>"
        )
        if len(review_rows) >= 20:
            break
    if len(data["review_queue"]) > len(review_rows):
        review_rows.append("<li>Xem toàn bộ mục rà soát trong <code>analysis/review_queue.json</code>.</li>")

    script = """
const $=(s,r=document)=>r.querySelector(s), $$=(s,r=document)=>[...r.querySelectorAll(s)];
async function copyText(value){try{await navigator.clipboard.writeText(value);return true}catch(e){return false}}
async function sharePage(){const data={title:document.title,text:'Biên bản cuộc họp',url:location.href};if(navigator.share){await navigator.share(data)}else{await copyText(location.href);alert('Đã sao chép liên kết')}}
async function shareSection(id){const url=location.href.split('#')[0]+'#'+id;await copyText(url);history.replaceState(null,'','#'+id);alert('Đã sao chép liên kết đến phần này')}
function filterSpeaker(name){$$('.turn').forEach(el=>el.classList.toggle('hidden',name!=='Tất cả'&&el.dataset.speaker!==name));$$('[data-speaker-filter]').forEach(b=>b.classList.toggle('active',b.dataset.speakerFilter===name))}
function sortTasks(key){const box=$('#task-list');$$('.task',box).sort((a,b)=>(a.dataset[key]||'').localeCompare(b.dataset[key]||'','vi')).forEach(el=>box.appendChild(el))}
function copyPlan(){copyText($$('#task-list .task').map(el=>el.innerText.trim()).join('\\n\\n')).then(()=>alert('Đã sao chép kế hoạch hành động'))}
$$('[data-share-section]').forEach(b=>b.addEventListener('click',()=>shareSection(b.dataset.shareSection)));
$$('[data-speaker-filter]').forEach(b=>b.addEventListener('click',()=>filterSpeaker(b.dataset.speakerFilter)));
"""
    doc = f"""<!doctype html>
<html lang="vi"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(data['title'])}</title><style>{css}</style></head><body>
<nav class="reader-bar"><strong>Biên bản cuộc họp</strong><a href="#muc-luc">Mục lục</a><button type="button" onclick="sharePage()">Chia sẻ</button><button type="button" onclick="window.print()">In / PDF</button></nav>
<main class="page">
<header class="hero"><div class="eyebrow">BIÊN BẢN CÓ DẪN CHỨNG</div><h1>{html.escape(data['title'])}</h1>
<div class="meta-grid"><div class="meta-card"><span>Ngày</span><strong>{html.escape(data['meeting_date'] or 'Chưa nêu')}</strong></div><div class="meta-card"><span>Thời lượng</span><strong>{duration_label}</strong></div><div class="meta-card"><span>Nguồn</span><strong>{html.escape(data['source']['type'])}</strong></div><div class="meta-card"><span>Trạng thái</span><strong>{'Cần rà soát' if data['review_queue'] else 'Đã kiểm tra'}</strong></div></div>
<div class="people">{participant_list}</div><div class="stats"><div class="stat"><b>{len(decisions)}</b><span>quyết định</span></div><div class="stat"><b>{len(actions)}</b><span>việc cần làm</span></div><div class="stat"><b>{len(data['review_queue'])}</b><span>mục cần rà soát</span></div><div class="stat"><b>{len(data['screenshots'])}</b><span>ảnh có mốc</span></div></div>
<div class="notice">Tài liệu chỉ phản ánh thông tin trong nguồn. Không dùng mục “Cần rà soát” để giao việc hoặc gửi email trước khi xác nhận.</div></header>
<nav class="toc" id="muc-luc"><strong>NỘI DUNG CHÍNH</strong><ol><li><a href="#y-chinh">Ý chính 30 giây</a></li><li><a href="#ke-hoach">Kế hoạch hành động</a></li><li><a href="#quyet-dinh">Quyết định</a></li><li><a href="#hinh-anh">Hình ảnh</a></li><li><a href="#dien-bien">Diễn biến</a></li><li><a href="#con-mo">Cần rà soát</a></li><li><a href="#kiem-chung">Kiểm chứng</a></li></ol></nav>
<section class="section" id="y-chinh"><div class="section-head"><span class="number">1</span><h2>Ý chính 30 giây</h2><button type="button" data-share-section="y-chinh">Chia sẻ phần</button></div><ul class="quick-list">{''.join(quick_items)}</ul><div class="summary-band"><span>{len(decisions)} quyết định</span><span>·</span><span>{len(actions)} việc cần làm</span><span>·</span><span>{len(data['review_queue'])} mục cần rà soát</span></div></section>
<section class="section" id="ke-hoach"><div class="section-head"><span class="number">2</span><h2>Kế hoạch hành động</h2><button type="button" data-share-section="ke-hoach">Chia sẻ phần</button></div><div class="controls"><button type="button" onclick="sortTasks('deadline')">Theo thời hạn</button><button type="button" onclick="sortTasks('owner')">Theo người</button><button type="button" onclick="copyPlan()">Sao chép kế hoạch</button></div><div class="cards" id="task-list">{''.join(action_cards)}</div></section>
<section class="section" id="quyet-dinh"><div class="section-head"><span class="number">3</span><h2>Quyết định đã chốt</h2><button type="button" data-share-section="quyet-dinh">Chia sẻ phần</button></div><div class="cards">{''.join(decision_cards)}</div></section>
<section class="section" id="hinh-anh"><div class="section-head"><span class="number">4</span><h2>Slide và hình ảnh</h2><button type="button" data-share-section="hinh-anh">Chia sẻ phần</button></div><div class="gallery">{figures or '<p class="muted">Không có screenshot từ nguồn. Với audio, chỉ dùng hình minh họa đã được gắn nhãn rõ ràng.</p>'}</div></section>
<section class="section" id="dien-bien"><div class="section-head"><span class="number">5</span><h2>Diễn biến cuộc họp</h2><button type="button" data-share-section="dien-bien">Chia sẻ phần</button></div><div class="controls"><button type="button" class="active" data-speaker-filter="Tất cả">Tất cả</button>{speaker_buttons}</div><div class="timeline">{''.join(timeline)}</div></section>
<section class="section" id="con-mo"><div class="section-head"><span class="number">6</span><h2>Câu hỏi còn mở và mục cần rà soát</h2><button type="button" data-share-section="con-mo">Chia sẻ phần</button></div><ul class="review-list">{''.join(review_rows) or '<li>Không có mục cần rà soát.</li>'}</ul></section>
<section class="section audit" id="kiem-chung"><div class="section-head"><span class="number">7</span><h2>Kiểm chứng</h2><button type="button" data-share-section="kiem-chung">Chia sẻ phần</button></div><p>Nguồn: {html.escape(data['source']['name'])} · SHA-256: <code>{html.escape(data['source']['sha256'])}</code></p><p>Tạo lúc {html.escape(data['generated_at'])}. Email chưa được gửi; việc xuất bản trang chia sẻ phải được người dùng phê duyệt.</p></section>
</main><script>{script}</script></body></html>"""
    share_target = output_dir / "bien-ban-share.html"
    share_target.write_text(doc, encoding="utf-8")
    (output_dir / "bien-ban.html").write_text(doc, encoding="utf-8")
    return share_target


def register_pdf_fonts() -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        (Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/arialbd.ttf")),
        (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")),
    ]
    for regular, bold in candidates:
        if regular.exists() and bold.exists():
            pdfmetrics.registerFont(TTFont("MeetingFont", str(regular)))
            pdfmetrics.registerFont(TTFont("MeetingFontBold", str(bold)))
            return "MeetingFont", "MeetingFontBold"
    return "Helvetica", "Helvetica-Bold"


def build_pdf(data: dict[str, Any], output_dir: Path) -> Path:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Image, LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError:
        fail("Thiếu reportlab. Hãy dùng Python bundled của Codex hoặc cài reportlab.")
    regular, bold = register_pdf_fonts()
    target = output_dir / "bien-ban.pdf"
    page_size = landscape(A4)
    doc = SimpleDocTemplate(str(target), pagesize=page_size, rightMargin=16 * mm, leftMargin=16 * mm, topMargin=15 * mm, bottomMargin=15 * mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("MeetingTitle", parent=styles["Title"], fontName=bold, fontSize=20, leading=24, textColor=colors.HexColor("#15394A"), spaceAfter=8)
    heading_style = ParagraphStyle("MeetingHeading", parent=styles["Heading2"], fontName=bold, fontSize=13, leading=16, textColor=colors.HexColor("#195B7A"), spaceBefore=13, spaceAfter=7)
    body_style = ParagraphStyle("MeetingBody", parent=styles["BodyText"], fontName=regular, fontSize=9.3, leading=13)
    small_style = ParagraphStyle("MeetingSmall", parent=body_style, fontSize=7.5, leading=10, textColor=colors.HexColor("#43515C"))
    header_style = ParagraphStyle("MeetingHeader", parent=body_style, fontName=bold, textColor=colors.white, alignment=TA_CENTER, fontSize=8)

    def p(value: Any, style=body_style) -> Paragraph:
        return Paragraph(html.escape(str(value)).replace("\n", "<br/>"), style)

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont(regular, 7)
        canvas.setFillColor(colors.HexColor("#66737F"))
        canvas.drawString(16 * mm, 9 * mm, data["title"][:90])
        canvas.drawRightString(page_size[0] - 16 * mm, 9 * mm, f"Trang {document.page}")
        canvas.restoreState()

    story: list[Any] = [p(data["title"], title_style)]
    story.append(p(f"Ngày họp: {data['meeting_date'] or 'Chưa nêu'} | Nguồn: {data['source']['name']}", small_style))
    story.append(p("Người tham gia: " + (", ".join(item["name"] for item in data["participants"]) or "Chưa cung cấp")))
    story.append(Spacer(1, 5 * mm))
    story.append(p("Tóm tắt kiểm chứng", heading_style))
    story.append(p(f"Ghi nhận {len(data['decisions'])} quyết định rõ ràng và {len(data['action_items'])} Action Item. Có {len(data['review_queue'])} vấn đề cần rà soát. Không dùng mục Cần rà soát để giao việc trước khi xác nhận."))
    story.append(p("Quyết định", heading_style))
    decision_rows = [[p("Mã", header_style), p("Quyết định", header_style), p("Người nói", header_style), p("Dẫn chứng", header_style), p("Trạng thái", header_style)]]
    for item in data["decisions"]:
        decision_rows.append([p(item["id"]), p(item["decision"]), p(item["speaker"]), p(item["evidence_ref"] + ": “" + item["evidence_quote"] + "”", small_style), p(item["review_status"])])
    if len(decision_rows) == 1:
        decision_rows.append([p(""), p("Không có quyết định rõ ràng trong nguồn."), p(""), p(""), p("")])
    decision_table = LongTable(decision_rows, colWidths=[15 * mm, 83 * mm, 35 * mm, 100 * mm, 28 * mm], repeatRows=1)
    decision_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#195B7A")),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D6DEE5")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F8FA")]),
    ]))
    story.append(decision_table)
    story.append(p("Action Items", heading_style))
    action_rows = [[p("Mã", header_style), p("Công việc", header_style), p("Người phụ trách", header_style), p("Deadline", header_style), p("Trạng thái", header_style), p("Dẫn chứng", header_style)]]
    for item in data["action_items"]:
        action_rows.append([p(item["id"]), p(item["task"]), p(item["owner"]), p(item["deadline_text"]), p(item["review_status"]), p(item["evidence_ref"] + ": “" + item["evidence_quote"] + "”", small_style)])
    if len(action_rows) == 1:
        action_rows.append([p(""), p("Không có Action Item rõ ràng trong nguồn."), p(""), p(""), p(""), p("")])
    action_table = LongTable(action_rows, colWidths=[14 * mm, 73 * mm, 34 * mm, 28 * mm, 29 * mm, 83 * mm], repeatRows=1)
    action_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#195B7A")),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D6DEE5")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F8FA")]),
    ]))
    story.append(action_table)
    story.append(p("Hàng đợi rà soát", heading_style))
    if data["review_queue"]:
        for item in data["review_queue"]:
            story.append(p(f"• {item['message']} ({item.get('evidence_ref', '')})"))
    else:
        story.append(p("Không có mục cần rà soát."))
    if data["screenshots"]:
        story.append(PageBreak())
        story.append(p("Ảnh tại các mốc quan trọng", heading_style))
        figure_cells: list[list[Any]] = []
        current_row: list[Any] = []
        for shot in data["screenshots"]:
            img_path = output_dir / shot["path"]
            image = Image(str(img_path))
            max_width, max_height = 68 * mm, 39 * mm
            ratio = min(max_width / image.imageWidth, max_height / image.imageHeight)
            image.drawWidth = image.imageWidth * ratio
            image.drawHeight = image.imageHeight * ratio
            image.hAlign = "CENTER"
            current_row.append([image, p(f"{shot['timestamp']} | {shot['caption']}", small_style)])
            if len(current_row) == 2:
                figure_cells.append(current_row)
                current_row = []
        if current_row:
            current_row.append("")
            figure_cells.append(current_row)
        figure_grid = Table(figure_cells, colWidths=[128 * mm, 128 * mm], hAlign="CENTER")
        figure_grid.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(figure_grid)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return target


def build_reader_pdf(data: dict[str, Any], output_dir: Path) -> Path:
    """Create the portrait, reader-facing PDF paired with bien-ban-share.html."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Image, KeepTogether, LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError:
        fail("Thiếu reportlab. Hãy dùng Python bundled của Codex hoặc cài reportlab.")

    regular, bold = register_pdf_fonts()
    target = output_dir / "bien-ban.pdf"
    doc = SimpleDocTemplate(
        str(target), pagesize=A4, rightMargin=15 * mm, leftMargin=15 * mm,
        topMargin=14 * mm, bottomMargin=15 * mm,
        title=f"{data['title']} - Biên bản", author="Meeting Minutes Agent",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ReaderTitle", parent=styles["Title"], fontName=bold, fontSize=22, leading=27, textColor=colors.HexColor("#252B30"), spaceAfter=7, alignment=0)
    eyebrow_style = ParagraphStyle("ReaderEyebrow", parent=styles["BodyText"], fontName=bold, fontSize=8, leading=10, textColor=colors.HexColor("#D85B25"), spaceAfter=5)
    heading_style = ParagraphStyle("ReaderHeading", parent=styles["Heading2"], fontName=bold, fontSize=13, leading=17, textColor=colors.HexColor("#165D72"), spaceBefore=10, spaceAfter=7)
    body_style = ParagraphStyle("ReaderBody", parent=styles["BodyText"], fontName=regular, fontSize=9, leading=13, textColor=colors.HexColor("#252B30"))
    body_bold = ParagraphStyle("ReaderBodyBold", parent=body_style, fontName=bold)
    small_style = ParagraphStyle("ReaderSmall", parent=body_style, fontSize=7.3, leading=9.5, textColor=colors.HexColor("#68737B"))
    quote_style = ParagraphStyle("ReaderQuote", parent=small_style, fontName=regular, fontSize=7.8, leading=10, leftIndent=7, textColor=colors.HexColor("#6C7378"))
    stat_value = ParagraphStyle("StatValue", parent=body_style, fontName=bold, fontSize=15, leading=16, textColor=colors.HexColor("#165D72"))
    stat_label = ParagraphStyle("StatLabel", parent=small_style, fontSize=7)
    header_style = ParagraphStyle("TableHeader", parent=small_style, fontName=bold, textColor=colors.HexColor("#5E6870"), fontSize=7.2)

    def p(value: Any, style=body_style) -> Paragraph:
        return Paragraph(html.escape(str(value)).replace("\n", "<br/>"), style)

    def footer(canvas, document):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#D7D4CE"))
        canvas.line(15 * mm, 11 * mm, A4[0] - 15 * mm, 11 * mm)
        canvas.setFont(regular, 7)
        canvas.setFillColor(colors.HexColor("#7A858D"))
        canvas.drawString(15 * mm, 7 * mm, data["title"][:76])
        canvas.drawRightString(A4[0] - 15 * mm, 7 * mm, f"Trang {document.page}")
        canvas.restoreState()

    story: list[Any] = [p("B I E N  B A N  C U O C  H O P", eyebrow_style), p(data["title"], title_style)]
    meta = [
        [p("Ngày", header_style), p(data["meeting_date"] or "Chưa nêu", body_bold), p("Nguồn", header_style), p(data["source"]["type"], body_bold)],
        [p("Thành phần", header_style), p(", ".join(x["name"] for x in data["participants"]) or "Chưa cung cấp", body_style), p("Tệp", header_style), p(data["source"]["name"], small_style)],
    ]
    meta_table = Table(meta, colWidths=[18 * mm, 65 * mm, 18 * mm, 79 * mm])
    meta_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 2)]))
    story.extend([meta_table, Spacer(1, 4 * mm)])

    stats = [[
        [p(len(data["decisions"]), stat_value), p("quyết định", stat_label)],
        [p(len(data["action_items"]), stat_value), p("công việc", stat_label)],
        [p(len(data["review_queue"]), stat_value), p("mục cần rà soát", stat_label)],
        [p(len(data["screenshots"]), stat_value), p("ảnh có mốc", stat_label)],
    ]]
    stat_table = Table(stats, colWidths=[45 * mm] * 4)
    stat_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F3F0EA")),
        ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#E2DED6")),
        ("INNERGRID", (0, 0), (-1, -1), 4, colors.white),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 9), ("RIGHTPADDING", (0, 0), (-1, -1), 9),
        ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.extend([stat_table, Spacer(1, 4 * mm)])

    story.append(p("1. Tóm tắt", heading_style))
    summary_items = []
    for item in data["decisions"][:3]:
        summary_items.append(f"Quyết định: {item['decision']}")
    for item in data["action_items"][:3]:
        summary_items.append(f"Việc cần làm: {item['task']} - {item['owner']}, hạn {item['deadline_text']}")
    if not summary_items:
        summary_items.append("Không có quyết định hoặc công việc rõ ràng trong nguồn.")
    for item in summary_items:
        story.append(p("- " + item))

    story.append(p("2. Diễn biến và dẫn chứng", heading_style))
    milestones = []
    for kind, items in (("Quyết định", data["decisions"]), ("Công việc", data["action_items"])):
        for item in items:
            milestones.append({
                "kind": kind,
                "second": item.get("timestamp_seconds"),
                "ref": item["evidence_ref"],
                "title": item.get("decision") or item.get("task"),
                "quote": item["evidence_quote"],
            })
    milestones.sort(key=lambda x: x["second"] if x["second"] is not None else 10**12)
    for item in milestones[:12]:
        block = Table([[p(item["ref"], body_bold), [p(item["kind"], body_bold), p(item["title"]), p('“' + item["quote"] + '”', quote_style)]]], colWidths=[22 * mm, 158 * mm])
        block.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LINEBELOW", (0, 0), (-1, -1), 0.35, colors.HexColor("#DDD8D0")),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.append(KeepTogether(block))

    if data["screenshots"]:
        story.extend([PageBreak(), p("3. Slide và hình ảnh", heading_style)])
        for shot in data["screenshots"][:6]:
            img_path = output_dir / shot["path"]
            image = Image(str(img_path))
            max_width, max_height = 82 * mm, 48 * mm
            ratio = min(max_width / image.imageWidth, max_height / image.imageHeight)
            image.drawWidth = image.imageWidth * ratio
            image.drawHeight = image.imageHeight * ratio
            caption = [p(f"{shot['timestamp']} - {shot['caption']}", body_bold), p("Ảnh chụp từ video tại mốc dẫn chứng.", small_style)]
            row = Table([[image, caption]], colWidths=[88 * mm, 92 * mm])
            row.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, -1), 0.35, colors.HexColor("#DDD8D0")), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
            story.append(KeepTogether(row))

    story.extend([PageBreak(), p("4. Quyết định", heading_style)])
    decision_rows = [[p("Mã", header_style), p("Quyết định", header_style), p("Người nói", header_style), p("Dẫn chứng", header_style)]]
    for item in data["decisions"]:
        decision_rows.append([p(item["id"], body_bold), p(item["decision"]), p(item["speaker"]), p(item["evidence_ref"] + " - " + item["review_status"], small_style)])
    if len(decision_rows) == 1:
        decision_rows.append([p(""), p("Không có quyết định rõ ràng trong nguồn."), p(""), p("")])
    decision_table = LongTable(decision_rows, colWidths=[14 * mm, 101 * mm, 32 * mm, 33 * mm], repeatRows=1)
    decision_table.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, 0), 1, colors.HexColor("#165D72")), ("LINEBELOW", (0, 1), (-1, -1), 0.35, colors.HexColor("#DDD8D0")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    story.append(decision_table)

    story.append(p("5. Kế hoạch hành động", heading_style))
    action_rows = [[p("Mã", header_style), p("Công việc", header_style), p("Phụ trách", header_style), p("Thời hạn", header_style), p("Căn cứ", header_style)]]
    for item in data["action_items"]:
        owner_style = ParagraphStyle("OwnerWarn", parent=body_style, fontName=bold, textColor=colors.HexColor("#B9382E")) if item["owner"] == "Chưa xác định" else body_style
        action_rows.append([p(item["id"], body_bold), p(item["task"]), p(item["owner"], owner_style), p(item["deadline_text"]), p(item["evidence_ref"], small_style)])
    if len(action_rows) == 1:
        action_rows.append([p(""), p("Không có Action Item rõ ràng trong nguồn."), p(""), p(""), p("")])
    action_table = LongTable(action_rows, colWidths=[14 * mm, 76 * mm, 32 * mm, 31 * mm, 27 * mm], repeatRows=1)
    action_table.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, 0), 1, colors.HexColor("#165D72")), ("LINEBELOW", (0, 1), (-1, -1), 0.35, colors.HexColor("#DDD8D0")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    story.append(action_table)

    story.append(p("6. Vấn đề còn mở và kiểm chứng", heading_style))
    unique_reviews = []
    seen = set()
    for item in data["review_queue"]:
        if item["message"] in seen:
            continue
        seen.add(item["message"])
        unique_reviews.append(item)
    if unique_reviews:
        for item in unique_reviews[:15]:
            story.append(p(f"- {item['message']} ({item.get('evidence_ref', '')})"))
    else:
        story.append(p("Không có mục cần rà soát."))
    story.extend([
        Spacer(1, 3 * mm),
        p(f"Kiểm chứng: {len(data['decisions'])} quyết định, {len(data['action_items'])} công việc, {len(data['review_queue'])} mục rà soát. Nguồn SHA-256: {data['source']['sha256']}", small_style),
        p("Email chưa được gửi. Bản web chỉ được xuất bản cho người khác xem sau khi người dùng xác nhận phạm vi chia sẻ.", small_style),
    ])
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return target


def write_email_drafts(data: dict[str, Any], output_dir: Path) -> list[Path]:
    drafts_dir = output_dir / "email-drafts"
    drafts_dir.mkdir(parents=True, exist_ok=True)
    grouped: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for item in data["action_items"]:
        if item["owner"] != "Chưa xác định" and item["owner_email"]:
            grouped[(item["owner"], item["owner_email"])].append(item)
    paths: list[Path] = []
    for (owner, email), items in grouped.items():
        msg = EmailMessage()
        msg["To"] = email
        msg["Subject"] = f"NHÁP - Công việc sau {data['title']}"
        msg["X-Unsent"] = "1"
        lines = [f"Chào {owner},", "", f"Dưới đây là các công việc được ghi nhận từ {data['title']}:", ""]
        for index, item in enumerate(items, start=1):
            lines.extend([
                f"{index}. {item['task']}",
                f"   Deadline: {item['deadline_text']}",
                f"   Dẫn chứng: {item['evidence_ref']} - “{item['evidence_quote']}”",
                "",
            ])
        lines.extend(["Vui lòng kiểm tra lại nội dung và deadline trước khi xác nhận.", "", "Trân trọng,"])
        msg.set_content("\n".join(lines), charset="utf-8")
        path = drafts_dir / f"nhap-{safe_name(owner)}.eml"
        path.write_bytes(msg.as_bytes())
        paths.append(path)
    unresolved = [item for item in data["action_items"] if item["owner"] == "Chưa xác định" or not item["owner_email"]]
    if unresolved:
        note = drafts_dir / "CAN-XAC-NHAN-NGUOI-NHAN.txt"
        note.write_text(
            "Không tạo email nháp cho các mục sau vì thiếu owner hoặc email đã xác nhận:\n\n"
            + "\n".join(f"- {item['id']}: {item['task']} ({'; '.join(item['review_reasons'])})" for item in unresolved),
            encoding="utf-8",
        )
        paths.append(note)
    return paths


def build_xlsx(data_path: Path, output_dir: Path) -> Path:
    node = shutil.which("node")
    if not node:
        fail("Không tìm thấy Node.js; không thể tạo action-items.xlsx bằng Artifact Tool.")
    script = SKILL_ROOT / "scripts" / "generate_xlsx.mjs"
    target = output_dir / "action-items.xlsx"
    preview = output_dir / "analysis" / "action-items-preview.png"
    command = [node, str(script), str(data_path), str(target), str(preview)]
    result = subprocess.run(command, cwd=str(SKILL_ROOT), capture_output=True, text=True, encoding="utf-8", errors="replace")
    if result.stdout.strip():
        print(result.stdout.strip())
    for _ in range(25):
        if target.exists() and target.stat().st_size > 0:
            break
        time.sleep(0.2)
    if not target.exists() or target.stat().st_size == 0:
        fail(f"Không tạo được Excel (exit={result.returncode}): {result.stderr.strip()}")
    try:
        with zipfile.ZipFile(target) as archive:
            if "xl/workbook.xml" not in archive.namelist():
                fail("File Excel đã tạo nhưng thiếu cấu trúc workbook.xml.")
    except zipfile.BadZipFile:
        fail("File Excel đã tạo không phải là XLSX hợp lệ.")
    inspect_sidecar = Path(str(target) + ".inspect.ndjson")
    if inspect_sidecar.exists():
        inspect_target = output_dir / "analysis" / "xlsx-inspect.ndjson"
        inspect_target.write_bytes(inspect_sidecar.read_bytes())
        inspect_sidecar.unlink()
    return target


def write_peas_report(data: dict[str, Any], output_dir: Path, warnings: list[str]) -> tuple[Path, Path]:
    unresolved = len(data["review_queue"])
    report = {
        "P": {
            "goal": "Tóm tắt chính xác; giữ quyết định, Action Items, owner, deadline và dẫn chứng; không bịa.",
            "observed": {
                "decisions": len(data["decisions"]),
                "action_items": len(data["action_items"]),
                "review_items": unresolved,
                "all_decisions_have_evidence": all(bool(item["evidence_ref"] and item["evidence_quote"]) for item in data["decisions"]),
                "all_actions_have_evidence": all(bool(item["evidence_ref"] and item["evidence_quote"]) for item in data["action_items"]),
            },
        },
        "E": {
            "local_source": data["source"]["path"],
            "google_drive": "Không truy cập trong lần chạy cục bộ này",
            "gmail": "Không truy cập; chỉ tạo file nháp .eml",
        },
        "A": {
            "read_transcript": True,
            "speaker_matching": True,
            "screenshots_created": len(data["screenshots"]),
            "share_html_created": (output_dir / "bien-ban-share.html").is_file(),
            "pdf_created": (output_dir / "bien-ban.pdf").is_file(),
            "web_published": False,
            "action_items_extracted": len(data["action_items"]),
            "email_drafts_created": len([item for item in data["action_items"] if item["owner"] != "Chưa xác định" and item["owner_email"]]),
            "email_sent": False,
        },
        "S": {
            "input_type": data["source"]["type"],
            "participants": len(data["participants"]),
            "segments": len(data["segments"]),
        },
        "warnings": warnings,
        "result": "PASS_WITH_REVIEW" if unresolved or warnings else "PASS",
    }
    json_path = output_dir / "analysis" / "peas-report.json"
    write_json(json_path, report)
    md_lines = [
        "# Báo cáo kiểm thử theo PEAS",
        "",
        f"Kết quả: **{report['result']}**",
        "",
        "## P Hiệu quả",
        "",
        f"- {len(data['decisions'])} quyết định; {len(data['action_items'])} Action Item; {unresolved} mục cần rà soát.",
        f"- Dẫn chứng quyết định đầy đủ: {report['P']['observed']['all_decisions_have_evidence']}.",
        f"- Dẫn chứng Action Item đầy đủ: {report['P']['observed']['all_actions_have_evidence']}.",
        "",
        "## E Môi trường",
        "",
        "- Đã xử lý tệp cục bộ. Google Drive và Gmail không được truy cập trong lần chạy này.",
        "",
        "## A Hành động",
        "",
        f"- Tạo bản web chia sẻ cục bộ, PDF A4, {len(data['screenshots'])} screenshot, {len(data['action_items'])} Action Item và email nháp cục bộ.",
        "- Chưa xuất bản web và không gửi email.",
        "",
        "## S Đầu vào",
        "",
        f"- Loại nguồn: {data['source']['type']}; {len(data['participants'])} người tham gia; {len(data['segments'])} lượt nói.",
    ]
    if warnings:
        md_lines.extend(["", "## Cảnh báo", "", *[f"- {warning}" for warning in warnings]])
    md_path = output_dir / "analysis" / "peas-report.md"
    md_path.write_text("\n".join(md_lines) + "\n", encoding="utf-8")
    return json_path, md_path


def write_manifest(output_dir: Path) -> Path:
    files = []
    for path in sorted(output_dir.rglob("*")):
        if path.is_file() and path.name != "manifest.json":
            files.append(
                {
                    "path": str(path.relative_to(output_dir)).replace("\\", "/"),
                    "bytes": path.stat().st_size,
                    "sha256": sha256_file(path),
                }
            )
    target = output_dir / "manifest.json"
    write_json(target, {"generated_at": datetime.now().astimezone().isoformat(), "files": files})
    return target


def run(args: argparse.Namespace) -> None:
    input_path = Path(args.input).expanduser().resolve()
    participants_path = Path(args.participants).expanduser().resolve()
    output_dir = Path(args.output).expanduser().resolve()
    if not input_path.is_file():
        fail(f"Không tìm thấy đầu vào: {input_path}")
    if not participants_path.is_file():
        fail(f"Không tìm thấy danh sách người tham gia: {participants_path}")
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "analysis").mkdir(exist_ok=True)
    config = read_json(Path(args.config).resolve() if args.config else DEFAULT_CONFIG)
    participants, alias_map = load_participants(participants_path)

    transcript_path: Path
    if args.transcript:
        transcript_path = Path(args.transcript).expanduser().resolve()
        if not transcript_path.is_file():
            fail(f"Không tìm thấy transcript: {transcript_path}")
    elif input_path.suffix.lower() in MEDIA_EXTS:
        transcript_path = transcribe_media(input_path, output_dir, args.whisper_model)
    else:
        transcript_path = input_path
    transcript_text = extract_text(transcript_path)
    normalized_copy = output_dir / "analysis" / "transcript-normalized.txt"
    normalized_copy.write_text(transcript_text.strip() + "\n", encoding="utf-8")
    segments, reviews = parse_segments(transcript_text, alias_map)
    decisions, actions = analyze(segments, participants, config, reviews)

    video: Path | None = None
    if args.video:
        video = Path(args.video).expanduser().resolve()
    elif input_path.suffix.lower() in VIDEO_EXTS:
        video = input_path
    if video and not video.is_file():
        fail(f"Không tìm thấy video: {video}")
    screenshots, warnings = capture_screenshots(video, decisions, actions, output_dir, int(config.get("screenshot_width", 960)))
    for warning in warnings:
        reviews.append({"type": "screenshot", "evidence_ref": "", "message": warning})

    source_kind = "media" if input_path.suffix.lower() in MEDIA_EXTS else input_path.suffix.lower().lstrip(".")
    data = {
        "title": args.title,
        "meeting_date": args.meeting_date,
        "generated_at": datetime.now().astimezone().isoformat(),
        "source": {"path": str(input_path), "name": input_path.name, "type": source_kind, "sha256": sha256_file(input_path)},
        "transcript_source": {"path": str(transcript_path), "sha256": sha256_file(transcript_path)},
        "participants": participants,
        "segments": segments,
        "decisions": decisions,
        "action_items": actions,
        "screenshots": screenshots,
        "review_queue": reviews,
        "safety": {"email_sent": False, "external_upload": False},
    }
    data_path = output_dir / "analysis" / "meeting_data.json"
    write_json(data_path, data)
    write_json(output_dir / "analysis" / "review_queue.json", reviews)
    build_html(data, output_dir)
    build_reader_pdf(data, output_dir)
    write_json(
        output_dir / "sharing.json",
        {
            "status": "local_only",
            "url": None,
            "audience": "local",
            "provider": None,
            "published_at": None,
            "note": "Bản web chưa được tải lên dịch vụ hosting. Cần người dùng xác nhận phạm vi chia sẻ trước khi xuất bản.",
        },
    )
    build_xlsx(data_path, output_dir)
    write_email_drafts(data, output_dir)
    write_peas_report(data, output_dir, warnings)
    write_manifest(output_dir)
    print(json.dumps({"status": "ok", "output": str(output_dir), "decisions": len(decisions), "actions": len(actions), "review_items": len(reviews)}, ensure_ascii=False))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Tạo biên bản họp có dẫn chứng, Action Items, Excel và email nháp.")
    parser.add_argument("--input", required=True, help="Transcript TXT/DOCX/PDF hoặc audio/video")
    parser.add_argument("--transcript", help="Transcript đi kèm khi --input là audio/video")
    parser.add_argument("--participants", required=True, help="JSON danh sách người tham gia")
    parser.add_argument("--video", help="Video dùng để chụp ảnh tại các mốc")
    parser.add_argument("--output", required=True, help="Thư mục đầu ra")
    parser.add_argument("--title", default="Biên bản cuộc họp", help="Tiêu đề biên bản")
    parser.add_argument("--meeting-date", default="", help="Ngày họp; chỉ dùng để hiển thị, không suy ra deadline")
    parser.add_argument("--config", help="JSON cấu hình; mặc định config/default.json")
    parser.add_argument("--whisper-model", default="small", help="Model faster-whisper cục bộ")
    return parser.parse_args()


if __name__ == "__main__":
    run(parse_args())
